"""Assemble the course submission folder Group_1 from the repository: the required documents, the dashboard guide and screenshots,
the scorecard and targeting tables, and a copy of the code. Markdown documents become PDFs through Microsoft Word, like the brief.

Run it last (after run-all, the release checks and committing the code), then commit Group_1: it is a copy and goes stale when
anything it copies changes. `--check` lists what is stale without rebuilding. Group_1/README.md is written by hand and kept.
"""
import argparse
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
from urllib.parse import unquote
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from telco_retention.common import ROOT, digest, now, read_json, remove_tree, write_json
from telco_retention.export import docx_to_pdf, markdown_to_docx, pdf_page_count
from telco_retention.html_dashboard import HTML_DASHBOARD

OUT = ROOT / "Group_1"
CODE = "09_Code"
# Folder -> {file name: repository source(s)}. Markdown sources are joined and converted to PDF; everything else is copied.
PACKAGE = {
    "01_Final_Report": {"Final_Report.pdf": "reports/final-report.pdf"},
    "02_Presentation": {"Presentation.pptx": "reports/presentation/retention-presentation.pptx",
                        "Presentation.pdf": "reports/presentation/retention-presentation.pdf",
                        "Speaker_Script.pdf": "reports/presentation/speaker-script.md",
                        "Demo_Script.pdf": "reports/presentation/demo-script.md"},
    "03_Executive_Brief": {"Executive_Brief.pdf": "reports/executive-brief.pdf", "Executive_Brief.docx": "reports/executive-brief.docx"},
    # Two editions: one HTML file that opens anywhere, and the Streamlit app, which runs from the code copy.
    "04_Dashboard": {"dashboard_html/Retention_Dashboard.html": HTML_DASHBOARD,
                     "dashboard_streamlit/Run_Dashboard.bat": "scripts/submission/Run_Dashboard.bat",
                     "dashboard_streamlit/Dashboard_User_Guide.pdf": "docs/user-guide.md",
                     **{f"dashboard_streamlit/screenshots/{p.name}": f"artifacts/qa/{p.name}" for p in sorted((ROOT / "artifacts/qa").glob("dashboard-*.png"))}},
    "05_Model_Card": {"Model_Card.pdf": "docs/model-card.md"},
    "06_Data_Dictionary_and_Source_Log": {"Data_Dictionary.pdf": "docs/data-dictionary.md",
                                          "Source_Log_and_Cleaning_Log.pdf": ["docs/source-log.md", "docs/cleaning-log.md"]},
    "07_Churn_Scorecard_and_Targeting_Table": {"Churn_Scorecard.pdf": "docs/churn-scorecard.md",
                                               "churn_scores_all_customers.csv": "reports/tables/scorecard.csv",
                                               "customer_targeting_table.csv": "reports/tables/targeting_recommended.csv"},
    "08_Retention_Playbook": {"Retention_Playbook.pdf": "docs/retention-playbook.md"},
}
# Line markers that pass through markdown_to_docx unchanged; tidy() styles the paragraph and removes them.
CODE_LINE, QUOTE_LINE = "§code§", "§quote§"
OS_FILES = {"desktop.ini", "thumbs.db", ".ds_store"}  # added to folders by Windows, file-sync tools and macOS (compared in lower case)
KEPT = {"README.md", "manifest.json"}  # the hand-written index and the build record; every other entry is a PACKAGE folder or CODE


def entries():
    for folder, files in PACKAGE.items():
        for name, sources in files.items():
            yield f"{folder}/{name}", [sources] if isinstance(sources, str) else sources


def files_in(folder):
    """Files under a folder, without those Windows and file-sync tools add to every folder."""
    return [p for p in folder.rglob("*") if p.is_file() and p.name.lower() not in OS_FILES] if folder.exists() else []


def unexpected_entries():
    """Top-level entries the build does not make, for example the folder of a deliverable that was renamed."""
    known = {*PACKAGE, CODE, *KEPT}
    return sorted(p.name for p in OUT.iterdir() if p.name not in known and p.name.lower() not in OS_FILES) if OUT.exists() else []


def provenance_problems(run_id):
    """The documents copied in must come from the run behind the tables and facts, not be PDFs kept from an earlier run,
    and must still be the files that export recorded."""
    deliverables = read_json(ROOT / "reports/deliverables-manifest.json")
    keys = ["run_id", "brief_pdf_run_id", "deck_pdf_run_id", "latex_run_id"]
    problems = [f"reports/deliverables-manifest.json: {k} is {deliverables.get(k)}, not run {run_id}; run export with Word, PowerPoint and pdflatex"
                for k in keys if deliverables.get(k) != run_id]
    exported = deliverables.get("sha256", {})
    for _, sources in entries():
        for source in sources:
            if source in exported and (not (ROOT / source).exists() or digest((ROOT / source).read_bytes()) != exported[source]):
                problems.append(f"{source} changed since export; run export again")
    return problems


def missing_tracked():
    """Tracked files deleted from the working tree: the copy of the code would silently lack them."""
    return [f"{n} is tracked but missing; restore it or commit its removal" for n in tracked_files() if not (ROOT / n).is_file()]


def git(*args):
    return subprocess.run(["git", *args], cwd=ROOT, capture_output=True, check=True).stdout.decode("utf-8")


def tracked_files():
    """The repository's tracked files, without the submission folder itself."""
    return [n for n in git("ls-files", "-z").split("\0") if n and not n.startswith(f"{OUT.name}/")]


def to_markdown(sources):
    """Join the sources and mark code blocks and quotes, which markdown_to_docx would otherwise print as Markdown syntax.
    Code passes through as numbered placeholders, so none of it is read as Markdown; tidy() puts the lines back verbatim."""
    lines, code_lines, code = [], [], False
    for source in sources:
        for line in (ROOT / source).read_text(encoding="utf-8").splitlines():
            if line.startswith("```"):
                code = not code
            elif code:
                lines.append(f"{CODE_LINE}{len(code_lines)}")
                code_lines.append(line)
            elif line.startswith(">"):
                lines.append(QUOTE_LINE + line[1:].strip() if line[1:].strip() else "")
            else:
                lines.append(line)
        lines.append("")
    return "\n".join(lines), code_lines


def tidy(docx_path, code_lines):
    """Code verbatim in a fixed-width font, quotes indented, and text left-aligned in tables (markdown_to_docx right-aligns every value cell for the brief's numbers)."""
    doc = Document(docx_path)
    for paragraph in doc.paragraphs:
        for marker in (CODE_LINE, QUOTE_LINE):
            if paragraph.runs and paragraph.runs[0].text.startswith(marker):
                paragraph.runs[0].text = paragraph.runs[0].text[len(marker):]
                paragraph.paragraph_format.left_indent = Cm(0.6)
                if marker == CODE_LINE:
                    paragraph.runs[0].text = code_lines[int(paragraph.runs[0].text)]
                    paragraph.runs[0].font.name, paragraph.runs[0].font.size = "Consolas", Pt(8.5)
    for table in doc.tables:
        for c in range(len(table.columns)):
            cells = [row.cells[c] for row in table.rows[1:]]
            wordy = [bool(re.search(r"[A-Za-z]{2,}", cell.text)) for cell in cells]  # "No", "Yes", "months"; not "3.7x" or "n/a"
            text_column = sum(wordy) * 2 > len(cells)  # mostly words: the whole column reads left to right
            for cell, words in zip(cells, wordy):
                if words or text_column:
                    for paragraph in cell.paragraphs:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.save(docx_path)


def readme_problems():
    readme = OUT / "README.md"
    if not readme.exists():
        return ["README.md is missing"]
    links = re.findall(r"\]\(([^)\s]+)(?:\s+\"[^\"]*\")?\)", readme.read_text(encoding="utf-8"))
    return [f"README.md links to a missing file: {t}" for t in links
            if not t.startswith(("http:", "https:", "#", "mailto:")) and not (OUT / unquote(t.split("#")[0])).exists()]


def build():
    run_id = read_json(ROOT / "reports/facts.json")["run"]["id"]
    blockers = provenance_problems(run_id) + missing_tracked() + [f"{OUT.name}/{n} is not made by this script: move or delete it first" for n in unexpected_entries()]
    if blockers:
        raise SystemExit("\n".join(["FAIL cannot build:", *blockers]))
    for folder in [*PACKAGE, CODE]:
        if (OUT / folder).exists():
            remove_tree(OUT / folder)
    documents = []
    with tempfile.TemporaryDirectory() as tmp:
        for file, sources in entries():
            target = OUT / file
            target.parent.mkdir(parents=True, exist_ok=True)
            if sources[0].endswith(".md"):
                md = Path(tmp) / f"{target.stem}.md"
                text, code_lines = to_markdown(sources)
                md.write_text(text, encoding="utf-8")
                docx = markdown_to_docx(md, Path(tmp) / f"{target.stem}.docx")
                tidy(docx, code_lines)
                docx_to_pdf(docx, target)
            else:
                assert len(sources) == 1, file
                shutil.copyfile(ROOT / sources[0], target)
            documents.append({"file": file, "sources": {s: digest((ROOT / s).read_bytes()) for s in sources}, "sha256": digest(target.read_bytes()),
                              **({"pages": pdf_page_count(target)} if target.suffix == ".pdf" else {})})
            print(f"  {file}", flush=True)
    code = tracked_files()
    for name in code:
        (OUT / CODE / name).parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(ROOT / name, OUT / CODE / name)
    print(f"  {CODE}/ ({len(code)} files)")
    write_json(OUT / "manifest.json", {"built_at_utc": now(), "run_id": run_id,
                                       "source_commit": git("rev-parse", "HEAD").strip(),
                                       "uncommitted_changes": bool(git("status", "--porcelain", "--", ".", f":(exclude){OUT.name}").strip()),
                                       "documents": documents, "code": {"folder": CODE, "files": len(code)}})
    problems = readme_problems()
    assert not problems, problems


def stale():
    """What differs from the repository: changed or missing copies, documents whose sources changed, stray files and folders,
    code that differs, and documents that do not come from the run behind the facts."""
    if not (OUT / "manifest.json").exists():
        return ["manifest.json is missing: build the folder"]
    manifest = read_json(OUT / "manifest.json")
    built = {d["file"]: d for d in manifest["documents"]}
    problems = provenance_problems(manifest["run_id"]) + missing_tracked() + [f"unexpected entry {n}" for n in unexpected_entries()]
    for file, sources in entries():
        entry, target = built.get(file), OUT / file
        if not entry or not target.exists() or digest(target.read_bytes()) != entry["sha256"] \
                or entry["sources"] != {s: digest((ROOT / s).read_bytes()) if (ROOT / s).exists() else None for s in sources}:
            problems.append(file)
    expected = {file for file, _ in entries()}
    problems += [f"stray file {f}" for f in sorted(p.relative_to(OUT).as_posix() for folder in PACKAGE for p in files_in(OUT / folder)) if f not in expected]
    code = OUT / CODE
    present = {p.relative_to(code).as_posix() for p in files_in(code)}
    tracked = set(tracked_files())
    problems += [f"{CODE}/{n}" for n in sorted(present ^ tracked)]
    problems += [f"{CODE}/{n}" for n in sorted(present & tracked) if not (ROOT / n).is_file() or (code / n).read_bytes() != (ROOT / n).read_bytes()]
    return problems + readme_problems()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--check", action="store_true", help="only report what is stale")
    if parser.parse_args().check:
        problems = stale()
        for problem in problems:
            print("STALE", problem)
        print("FAIL submission folder is stale: run scripts/build_submission.py" if problems else "PASS submission folder matches the repository")
        sys.exit(1 if problems else 0)
    print(f"Building {OUT.name}/")
    build()
    print("PASS", OUT.name, "built")
