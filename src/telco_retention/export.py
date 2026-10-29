"""Produce the submission files: executive brief (DOCX + PDF), presentation (PPTX + PDF, script), the LaTeX final report (PDF)
and the single-file HTML dashboard."""
from __future__ import annotations

import re
from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from .common import ROOT, digest, latest_run, now, read_json, write_json

INK, SECONDARY, ACCENT = "0B0B0B", "52514E", "2A78D6"
IMAGE = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)")


class ExportUnavailable(RuntimeError):
    pass


def _runs(paragraph, text, size=None, color=INK):
    """Add text with **bold** and *italic* spans; Markdown links keep only their text."""
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    for part in re.split(r"(\*\*.+?\*\*|(?<![*\w])\*[^*\s][^*]*\*(?![*\w]))", text):
        if not part:
            continue
        bold, italic = part.startswith("**") and part.endswith("**") and len(part) > 4, False
        if bold:
            part = part[2:-2]
        elif len(part) > 2 and part.startswith("*") and part.endswith("*"):
            part, italic = part[1:-1], True
        run = paragraph.add_run(part.replace("`", ""))
        run.bold, run.italic = bold, italic
        run.font.color.rgb = RGBColor.from_string(color)
        if size:
            run.font.size = Pt(size)


def _table_rows(lines):
    rows = [[c.strip() for c in line.strip().strip("|").split("|")] for line in lines]
    return [r for r in rows if not all(set(c) <= set("-: ") for c in r)]


def markdown_to_docx(md_path, out_path, image_width_cm=15.5):
    md_path, out_path = Path(md_path), Path(out_path)
    doc = Document()
    section = doc.sections[0]
    section.page_height, section.page_width = Cm(29.7), Cm(21.0)
    section.left_margin = section.right_margin = Cm(1.9)
    section.top_margin = section.bottom_margin = Cm(1.6)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.08
    for level, size in [(1, 16), (2, 12), (3, 10.5)]:
        style = doc.styles[f"Heading {level}"]
        style.font.name, style.font.size, style.font.bold = "Calibri", Pt(size), True
        style.font.color.rgb = RGBColor.from_string(INK if level != 2 else ACCENT)
        style.paragraph_format.space_before, style.paragraph_format.space_after = Pt(8 if level > 1 else 0), Pt(3)
    text = re.sub(r"<!--.*?-->", "", md_path.read_text(encoding="utf-8"), flags=re.S)
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            doc.add_heading(line.lstrip("#").strip(), level=min(level, 3))
        elif line.lstrip().startswith("|"):
            block = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                block.append(lines[i])
                i += 1
            rows = _table_rows(block)
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Light Grid Accent 1"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    cell = table.cell(r, c)
                    cell.text = ""
                    _runs(cell.paragraphs[0], value if r else f"**{value}**", size=9)
                    if c and r:
                        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
            continue
        elif IMAGE.search(line):
            alt, src = IMAGE.search(line).groups()
            doc.add_picture(str((md_path.parent / src).resolve()), width=Cm(image_width_cm))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif re.match(r"^\s*\d+\.\s+", line):
            # Keep the Markdown number: Word's list styles would continue numbering across separate lists.
            number, rest = re.match(r"^\s*(\d+)\.\s+(.*)$", line).groups()
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent, paragraph.paragraph_format.first_line_indent = Cm(0.63), Cm(-0.63)
            paragraph.paragraph_format.space_after = Pt(1)
            _runs(paragraph, f"{number}.\t{rest}")
            paragraph.paragraph_format.tab_stops.add_tab_stop(Cm(0.63))
        elif re.match(r"^\s*[-*]\s+", line):
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(1)
            _runs(paragraph, re.sub(r"^\s*[-*]\s+", "", line))
        else:
            paragraph = doc.add_paragraph()
            _runs(paragraph, line.strip(), color=SECONDARY if line.startswith("_") else INK)
        i += 1
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path


def docx_to_pdf(docx_path, pdf_path):
    try:
        import pythoncom
        from docx2pdf import convert
        pythoncom.CoInitialize()
        convert(str(docx_path), str(pdf_path))
    except Exception as exc:  # Word missing or COM failure
        raise ExportUnavailable(f"PDF export needs Microsoft Word (docx2pdf): {exc}") from exc
    if not Path(pdf_path).exists():
        raise ExportUnavailable("Word did not produce a PDF")
    return Path(pdf_path)


def pptx_to_pdf(pptx_path, pdf_path):
    try:
        import pythoncom
        import win32com.client
        pythoncom.CoInitialize()
        app = win32com.client.Dispatch("PowerPoint.Application")
        deck = app.Presentations.Open(str(Path(pptx_path).resolve()), WithWindow=False)
        deck.SaveAs(str(Path(pdf_path).resolve()), 32)
        deck.Close()
    except Exception as exc:
        raise ExportUnavailable(f"Deck PDF export needs Microsoft PowerPoint: {exc}") from exc
    return Path(pdf_path)


def pdf_page_count(pdf_path):
    from pypdf import PdfReader
    return len(PdfReader(str(pdf_path)).pages)


def _check_generated(path, run):
    """Generated sources name their run in the first line; refuse to convert one left over from another run."""
    first = Path(path).read_text(encoding="utf-8").partition("\n")[0] if Path(path).exists() else ""
    if f"for run {run.name}." not in first:
        raise ValueError(f"{Path(path).name} is not from run {run.name}; run `report` for it first")


def _unchanged(path, previous, root=ROOT):
    """True while a file is exactly the one the previous export hashed under its path."""
    path = Path(path)
    return path.exists() and previous.get("sha256", {}).get(path.relative_to(root).as_posix()) == digest(path.read_bytes())


def convert_pdfs(jobs, mode, run_id, previous, root=ROOT):
    """Convert each job (name, converter, source, target, pages key) and record every PDF's run separately.

    A PDF that cannot be rebuilt keeps its previous page count and run only while it is still the file the previous
    export hashed at that path; a missing, replaced or half-written PDF is not described."""
    out, errors = {}, []
    for name, convert, source, target, pages in jobs:
        if mode != "skip":
            try:
                convert(source, target)
                out.update({pages: pdf_page_count(target), f"{name}_pdf_run_id": run_id})
                continue
            except ExportUnavailable as exc:
                if mode == "require":
                    raise
                errors.append(str(exc))
        if pages in previous and _unchanged(target, previous, root):
            # Older manifests recorded one run for both PDFs, or only the export's run.
            out.update({pages: previous[pages], f"{name}_pdf_run_id": previous.get(f"{name}_pdf_run_id", previous.get("pdf_run_id", previous.get("run_id")))})
    out["pdf_status"] = "skipped" if mode == "skip" else "unavailable" if len(errors) == len(jobs) else "partial" if errors else "current"
    if errors:
        out["pdf_error"] = "; ".join(errors)
    return out


def export(run=None, pdf="auto", latex="auto"):
    """pdf: Word/PowerPoint conversion of the brief and deck; latex: pdflatex build of the final report. Each is auto, require or skip.
    A missing tool is skipped under auto; a LaTeX compile error is always raised."""
    from .html_dashboard import build_html_dashboard
    from .latex import LATEX_DIR, LATEX_DOCUMENTS, LatexUnavailable, compile_latex
    from .presentation import SPEAKERS, build_presentation, lint
    run = run or latest_run()
    reports = ROOT / "reports"
    facts = read_json(reports / "facts.json")
    if facts["run"]["id"] != run.name:
        raise ValueError(f"reports/facts.json is from run {facts['run']['id']}; run `report` for {run.name} first")
    brief_md, brief_docx, brief_pdf = reports / "executive-brief.md", reports / "executive-brief.docx", reports / "executive-brief.pdf"
    deck_pptx, deck_pdf = reports / "presentation/retention-presentation.pptx", reports / "presentation/retention-presentation.pdf"
    script_md = reports / "presentation/speaker-script.md"
    report_pdf = ROOT / LATEX_DOCUMENTS["latex/report.tex.j2"][1]
    _check_generated(brief_md, run)
    markdown_to_docx(brief_md, brief_docx, image_width_cm=15)
    script = build_presentation(facts, deck_pptx, script_md)
    dashboard_html = build_html_dashboard(run)
    timed = [e for e in script if e["speaker"] in SPEAKERS]
    previous_path = reports / "deliverables-manifest.json"
    previous = read_json(previous_path) if previous_path.exists() else {}
    manifest = {"run_id": run.name, "created_at_utc": now(), "brief_words": len(re.sub(r"<!--.*?-->", "", brief_md.read_text(encoding="utf-8"), flags=re.S).split()),
                "slides": len(script), "slides_with_notes": sum(bool(e["text"].strip()) for e in script), "talk_seconds": sum(e["seconds"] for e in timed),
                "speaker_seconds": {s: sum(e["seconds"] for e in timed if e["speaker"] == s) for s in SPEAKERS}, "deck_lint": lint(script),
                "latex_status": "skipped"}
    manifest.update(convert_pdfs([("brief", docx_to_pdf, brief_docx, brief_pdf, "brief_pages"), ("deck", pptx_to_pdf, deck_pptx, deck_pdf, "deck_pdf_pages")],
                                 pdf, run.name, previous))
    if latex != "skip":
        for source, _ in LATEX_DOCUMENTS.values():
            _check_generated(ROOT / LATEX_DIR / source, run)
        try:
            manifest["report_pages"] = compile_latex()[LATEX_DOCUMENTS["latex/report.tex.j2"][1]]
            manifest.update(latex_status="current", latex_run_id=run.name)
        except LatexUnavailable as exc:
            if latex == "require":
                raise
            manifest.update(latex_status="unavailable", latex_error=str(exc)[-1000:])
    if manifest["latex_status"] != "current" and _unchanged(report_pdf, previous):
        for key in ["report_pages", "latex_run_id"]:
            if key in previous:
                manifest[key] = previous[key]
    manifest["sha256"] = {str(p.relative_to(ROOT)).replace("\\", "/"): digest(p.read_bytes()) for p in [brief_docx, brief_pdf, deck_pptx, deck_pdf, script_md, report_pdf, dashboard_html] if p.exists()}
    write_json(reports / "deliverables-manifest.json", manifest)
    return manifest
