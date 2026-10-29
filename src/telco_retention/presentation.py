"""Presentation deck on the HUST template (PPTX), with each speaker's script in the slide notes, built from the run's facts."""
from __future__ import annotations

import io
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from .common import ROOT
from .templating import TEMPLATES, num, pct, usd

TEMPLATE = TEMPLATES / "presentation/hust-template.pptx"  # HUST seminar template with its slides removed; masters only
LOGO = TEMPLATES / "presentation/hust-logo.png"
DASHBOARD = ROOT / "artifacts/qa/dashboard-decision.png"
FIGURES = ROOT / "reports/figures"
NAVY, RED, CARD, INK, MUTED, WHITE = (RGBColor.from_string(c) for c in ["0B2E63", "BE1E2D", "EEF2F8", "1F2937", "6B7280", "FFFFFF"])
FONT = "Calibri"
MASTERS = {"title": 0, "outline": 1, "content": 2, "thanks": 3, "closing": 4}
SPEAKERS = ["Kiet", "Hieu", "Dung"]  # timed talk; Phong runs the live demo after it
DEMO_SPEAKER = "Phong"
TALK_SECONDS = 480
SECTIONS = [("Business problem", "Decision, data, risk versus value"),
            ("Model and decision rule", "Calibrated churn model, expected value, optimiser"),
            ("Results and robustness", "Same-budget comparison, budget size, sensitivity"),
            ("Recommendation and demo", "Randomised pilot, live dashboard")]


# ---------------------------------------------------------------- drawing helpers
def _font(run, size, color=INK, bold=False, italic=False):
    font = run.font
    font.name, font.size, font.bold, font.italic = FONT, Pt(size), bold, italic
    font.color.rgb = color


def _bullet(paragraph, size):
    """Hanging-indent bullet, as in the template's text slides."""
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(Inches(0.25)))
    pPr.set("indent", str(-Inches(0.25)))
    clr = pPr.makeelement(qn("a:buClr"), {})
    clr.append(clr.makeelement(qn("a:srgbClr"), {"val": "0B2E63"}))
    pPr.append(clr)
    pPr.append(pPr.makeelement(qn("a:buSzPts"), {"val": str(int(size * 100))}))
    pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
    pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•"}))


def _fill(frame, paragraphs, size, color, bold, align, bullets=False, space_after=6):
    """paragraphs: strings, or lists of (text, style) runs where style may set size, color, bold, italic."""
    for i, para in enumerate(paragraphs):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        for text, style in ([(para, {})] if isinstance(para, str) else para):
            run = p.add_run()
            run.text = text
            _font(run, style.get("size", size), style.get("color", color), style.get("bold", bold), style.get("italic", False))
        if bullets:
            _bullet(p, size)


def text(slide, x, y, w, h, paragraphs, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, bullets=False, space_after=6):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap, frame.vertical_anchor = True, anchor
    _fill(frame, paragraphs, size, color, bold, align, bullets, space_after)
    return box


def card(slide, x, y, w, h, paragraphs=(), fill=CARD, size=16, color=INK, bold=False, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, bullets=False, space_after=4):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.adjustments[0] = 0.08
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    shape.shadow.inherit = False
    frame = shape.text_frame
    frame.word_wrap, frame.vertical_anchor = True, anchor
    frame.margin_left = frame.margin_right = Inches(0.18)
    frame.margin_top = frame.margin_bottom = Inches(0.08)
    _fill(frame, paragraphs, size, color, bold, align, bullets, space_after)
    return shape


def stat(slide, x, y, w, h, value, caption, value_size=32, caption_size=13):
    """Big red number over a short caption: the template's callout card."""
    return card(slide, x, y, w, h, [[(value, {"size": value_size, "bold": True, "color": RED})], [(caption, {"size": caption_size})]], space_after=2)


def banner(slide, x, y, w, h, paragraphs, size=20):
    return card(slide, x, y, w, h, paragraphs, fill=NAVY, size=size, color=WHITE, bold=True)


def operator(slide, x, y, w, h, symbol):
    text(slide, x, y, w, h, [symbol], size=28, color=RED, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def title(slide, value):
    text(slide, 0.37, 0.12, 12.59, 0.48, [value], size=28, color=WHITE, bold=True)


def page_number(slide, n):
    text(slide, 12.08, 7.08, 0.9, 0.33, [str(n)], size=11, color=MUTED, align=PP_ALIGN.RIGHT)


def caption(slide, x, y, w, value, h=0.35):
    text(slide, x, y, w, h, [value], size=12, color=MUTED)


def picture(slide, source, x, y, w=None, h=None):
    """Picture scaled to the given width or height, keeping its aspect ratio."""
    with Image.open(source) as image:
        ratio = image.width / image.height
    if hasattr(source, "seek"):
        source.seek(0)
    w, h = (w, w / ratio) if w is not None else (h * ratio, h)
    return slide.shapes.add_picture(source if hasattr(source, "seek") else str(source), Inches(x), Inches(y), Inches(w), Inches(h))


def table(slide, x, y, widths, rows, row_h=0.5, size=14, highlight=None, text_columns=1):
    """Navy header, alternating light rows; the highlighted row is bold with red numbers, like the template.
    The first `text_columns` columns are left-aligned, the rest (numbers) right-aligned."""
    shape = slide.shapes.add_table(len(rows), len(widths), Inches(x), Inches(y), Inches(sum(widths)), Inches(row_h * len(rows)))
    grid = shape.table
    for j, width in enumerate(widths):
        grid.columns[j].width = Inches(width)
    for i, row in enumerate(rows):
        grid.rows[i].height = Inches(row_h)
        for j, value in enumerate(row):
            cell = grid.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if i == 0 else CARD if i % 2 else WHITE
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            frame = cell.text_frame
            frame.word_wrap = True
            p = frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j < text_columns else PP_ALIGN.RIGHT
            run = p.add_run()
            run.text = str(value)
            if i == 0:
                _font(run, size + 1, WHITE, True)
            elif i == highlight:
                _font(run, size, NAVY if j == 0 else RED, True)
            else:
                _font(run, size, INK)
    return shape


def share(p):
    """Whole percent, except near-certain shares, which rounding would turn into a claim of 100%."""
    return f"{p * 100:.2f}%" if 0.995 <= p < 1 else pct(p, 0)


def crop_dashboard(path, height_px=1000):
    """Top of the full-page dashboard screenshot: header, KPI cards and the recommendation."""
    image = Image.open(path).convert("RGB")
    stream = io.BytesIO()
    image.crop((0, 0, image.width, min(height_px, image.height))).save(stream, "PNG")
    stream.seek(0)
    return stream


# ---------------------------------------------------------------- deck
class Deck:
    def __init__(self, template=TEMPLATE):
        self.prs = Presentation(str(template))
        self.script = []
        self.clock = 0

    def slide(self, kind, heading, speaker, seconds, notes):
        layout = self.prs.slide_masters[MASTERS[kind]].slide_layouts[0]
        slide = self.prs.slides.add_slide(layout)
        n = len(self.prs.slides)
        if speaker in SPEAKERS:
            start, self.clock = self.clock, self.clock + seconds
            cue = f"[{speaker} | {start // 60}:{start % 60:02d}-{self.clock // 60}:{self.clock % 60:02d} | {seconds} s]"
        elif speaker == DEMO_SPEAKER:
            cue = f"[{speaker} | live demo before the closing words, not part of the {TALK_SECONDS // 60}-minute talk | full timed steps in demo-script.md]"
        else:
            cue = "[No speaker]"
        slide.notes_slide.notes_text_frame.text = f"{cue}\n{notes}"
        self.script.append({"slide": n, "title": heading, "speaker": speaker, "seconds": seconds, "cue": cue, "text": notes})
        if kind == "content":
            title(slide, heading)
            page_number(slide, n)
        return slide

    def outline(self, speaker, seconds, notes, active=None):
        s = self.slide("outline", "Outline", speaker, seconds, notes)
        text(s, 4.99, 0.44, 8.08, 0.48, ["Outline"], size=28, color=NAVY, bold=True)
        for i, (name, detail) in enumerate(SECTIONS):
            on = active is None or i == active
            numeral = ["I.", "II.", "III.", "IV."][i] + "  "
            runs = [[(numeral, {"size": 22, "bold": True, "color": RED if on else MUTED}), (name, {"size": 22, "bold": on, "color": NAVY if on else MUTED})],
                    [(detail, {"size": 14, "color": INK if on else MUTED})]]
            box = card(s, 4.98, 1.3 + i * 1.07, 7.9, 0.94, runs, fill=CARD, align=PP_ALIGN.LEFT, space_after=0)
            box.text_frame.margin_left = Inches(0.25)
            if not on:
                box.fill.background()
        return s

    def save(self, path):
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        props = self.prs.core_properties
        props.title, props.author, props.subject = "Profit-Aware Telecom Churn & Retention Targeting", "Group 1", "Business Analytics capstone, Topic 7"
        self.prs.save(str(path))
        return path


def build_presentation(facts, out_path, script_path=None):
    """Write the deck and, optionally, the rehearsal script; returns the script entries (one per slide)."""
    f = facts
    c, q = f["comparison"], f["default_offer"]["save_rate"]
    chosen, hl, m, reco, pilot = c[f["strategy"]], f["headline"], f["model"], f["recommended"], f["pilot"]
    mc, quad = f["monte_carlo"]["by"], f["quadrant"]
    rbe = pilot["break_even_save_rate"]  # on the scale the pilot measures: the list's mix save rate
    rbe_standard = reco["break_even"]["break_even_default_save_rate"]  # the same break-even as the standard offer's save rate
    q_mix = pilot["assumed_save_rate"]
    senior = f["fairness_by"][f"{f['strategy']}|SeniorCitizen|1"]
    offers = {o["offer_id"]: o for o in f["offers"]}
    deck = Deck()

    # 1. Title
    s = deck.slide("title", "Title", "Kiet", 15,
                   "Good morning, Professor, and everyone. We are Group 1, presenting Topic 7: profit-aware telecom churn and retention targeting. "
                   "Hieu, Dung and I will present, and Phong will demo our dashboard at the end.")
    s.shapes.add_picture(str(LOGO), Inches(0.43), Inches(0.3), Inches(3.47), Inches(1.26))
    text(s, 0.76, 2.15, 9.0, 1.45, ["Profit-Aware Telecom Churn", "& Retention Targeting"], size=36, color=NAVY, bold=True, space_after=0)
    text(s, 0.76, 3.7, 9.0, 0.42, ["Introduction to Business Analytics · Capstone Topic 7 · Assoc. Prof. Nguyen Binh Minh"], size=14, color=INK)
    text(s, 0.76, 4.2, 5.6, 1.7, [[("Group 1", {"bold": True, "color": NAVY})], "Pham Tuan Kiet · 20252751M", "Nguyen Trung Hieu · 20261059M",
                                  "Bui Tien Dung · 20261204M", "Pham Hung Phong · 20261061M"], size=15, space_after=1)

    # 2. Outline
    deck.outline("Kiet", 10, "We follow CRISP-DM: the business problem and the data, then the model and the decision rule, "
                             "then results and robustness, and finally our recommendation.")

    # 3. The decision
    var_scaled = f["value_at_risk"]["expected_scaled"]
    s = deck.slide("content", "The Decision: Who Gets a Retention Offer?", "Kiet", 50,
                   f"In our sample company {pct(f['data']['churn_rate'])} of customers left, which puts about {usd(var_scaled)} "
                   "of twelve-month margin at risk. The retention team has a fixed budget, and today it calls the customers with the highest churn risk. "
                   "But every offer costs money, and a risky customer with a small bill may not be worth saving. So our question is: with a limited budget, "
                   "who should receive an offer when both risk and value count? Our answer is to score each customer and offer by expected value: "
                   "churn probability, times the save rate, the share of would-be leavers an offer keeps, times customer value, minus the cost. "
                   "We contact a customer only when it is positive.")
    banner(s, 0.5, 1.35, 12.35, 0.95, ["Given a limited retention budget, which customers should receive an offer when both churn risk and expected value count?"], size=20)
    text(s, 0.5, 2.55, 12.35, 1.6, [f"{pct(f['data']['churn_rate'])} of customers left; churn puts about {usd(var_scaled)} of 12-month margin at risk",
                                    "Today the fixed budget goes to the customers most likely to leave",
                                    "But every offer costs money, and not every risky customer is worth saving"], size=18, bullets=True)
    x = 0.6
    for i, (head, sub) in enumerate([("Churn risk", "p, from the model"), ("Save rate", "q, assumed per offer"), ("Customer value", "V, 12-month margin"),
                                     ("Cost", "offer + contact"), ("Offer if EV > 0", "EV = p × q × V − cost")]):
        last = i == 4
        card(s, x, 4.45, 2.05, 1.35, [[(head, {"size": 17, "bold": True, "color": WHITE if last else NAVY})], [(sub, {"size": 12, "color": WHITE if last else INK})]],
             fill=RED if last else CARD, space_after=2)
        if not last:
            operator(s, x + 2.05, 4.45, 0.45, 1.35, ["×", "×", "−", "→"][i])
        x += 2.5
    caption(s, 0.5, 6.05, 12.35, "Save rate = share of would-be leavers an offer keeps. Money in USD of the fictional sample company.")

    # 4. Data and validation design
    quality = f["data"]["quality"]
    s = deck.slide("content", "Data and Validation Design", "Kiet", 45,
                   f"We use the IBM Telco Customer Churn sample from Kaggle: {num(f['data']['rows'])} customers of a fictional US operator. "
                   "The unit of analysis is the customer, the target is churn, and the decision is whether to make an offer, and which one, within the budget. "
                   f"Quality checks found no duplicates and only {quality['total_charges_filled_zero']} blank total charges, all for brand-new customers, which we set to zero. "
                   "The data is one snapshot without dates, so we used a stratified 60/20/20 split: every choice is made on validation, and the test set is used once. "
                   "To avoid leakage we excluded churn reasons and IBM's churn score, which are known only after a customer leaves.")
    stat(s, 0.5, 1.4, 3.7, 1.4, num(f["data"]["rows"]), "customers, IBM Telco sample (fictional US operator)")
    stat(s, 0.5, 3.0, 3.7, 1.4, pct(f["data"]["churn_rate"]), "churned in the sample window")
    stat(s, 0.5, 4.6, 3.7, 1.4, "60 / 20 / 20", "train / validation / test, stratified")
    splits = f["data"]["splits"]
    table(s, 4.6, 1.4, [2.45, 5.8], [
        ["Item", "Decision"],
        ["Unit of analysis", "One customer"],
        ["Target", "Churn: left during the sample window (yes / no)"],
        ["Decision", "Offer or not, and which offer, within the budget"],
        ["Data quality", f"{quality['duplicates_removed']} duplicates; {quality['total_charges_filled_zero']} blank TotalCharges (tenure 0) set to 0"],
        ["Leakage control", "Churn reasons and IBM score excluded; TotalCharges dropped"],
        ["Validation", f"{num(splits['train'])} / {num(splits['validation'])} / {num(splits['test'])}; test used once"],
        ["Fairness", "Gender and senior status audited, never model inputs"],
    ], row_h=0.575, size=14, text_columns=2)

    # 5. Risk is not value
    hrlv, hrhv = quad["hrlv"], quad["hrhv"]
    eda = f["eda"]
    s = deck.slide("content", "Key Insight: Churn Risk Is Not Customer Value", "Kiet", 40,
                   "The exploratory analysis gave us our key insight. Churn is highest for month-to-month contracts, new customers, fibre internet and electronic-check payers. "
                   f"But risk and value only partly overlap. The high-risk, low-value group holds {pct(hrlv['value_at_risk_share'], 0)} of the value at risk, "
                   f"yet each of these customers is worth little. Only {hrhv['customers']} test customers are both high-risk and high-value, and today's risk-only rule "
                   f"reaches just {pct(c['risk']['high_risk_high_value_coverage'], 0)} of them. So ranking by risk alone spends the budget in the wrong place. "
                   "Hieu will now show how we fix that.")
    picture(s, FIGURES / "risk_value.png", 0.5, 1.3, h=5.15)
    stat(s, 8.3, 1.4, 4.55, 1.5, pct(hrlv["value_at_risk_share"], 0), "of value at risk sits with high-risk, low-value customers, each worth little")
    stat(s, 8.3, 3.1, 4.55, 1.5, str(hrhv["customers"]), f"test customers are high-risk and high-value; risk-only targeting reaches {pct(c['risk']['high_risk_high_value_coverage'], 0)}")
    text(s, 8.3, 4.8, 4.55, 1.6, [f"Churn: month-to-month {pct(eda['Contract'][0]['churn_rate'], 0)} vs two-year {pct(eda['Contract'][-1]['churn_rate'], 0)}",
                                  f"New customers {pct(eda['tenure_segment'][0]['churn_rate'], 0)}, fibre {pct(eda['InternetService'][0]['churn_rate'], 0)}, "
                                  f"electronic check {pct(eda['PaymentMethod'][0]['churn_rate'], 0)}",
                                  "Associations, not causes"], size=14, bullets=True, space_after=4)

    # 6. Outline: model and decision rule
    deck.outline("Hieu", 5, "Thank you, Kiet. I will start with the churn model.", active=1)

    # 7. Churn model
    test = {row["model"]: row for row in m["test"]}
    odds = {row["term"]: row["odds_ratio"] for row in m["odds_ratios"]}
    sel = m["selected_test"]
    s = deck.slide("content", "Churn Model: Simple, Calibrated, Good Enough", "Hieu", 50,
                   "We compared three model families: logistic regression as the baseline, a random forest as the required tree model, and XGBoost. "
                   "Each was tuned with five-fold cross-validation and calibrated. On the untouched test set they rank customers almost equally well, so, "
                   "by a rule we fixed in advance, we deploy the simplest: calibrated logistic regression. "
                   f"Its ROC-AUC is {num(sel['roc_auc'], 2)} and its PR-AUC {num(sel['average_precision'], 2)}, against {num(f['data']['churn_rate'], 2)} for a model that knows nothing. "
                   f"Calibration is essential here, because we multiply the probability by money: the calibration error is only {num(sel['ece'], 3)}. "
                   f"The strongest drivers are month-to-month contracts, which multiply the odds of churn by about {num(odds['Contract_Month-to-month'], 1)}, "
                   "and fibre internet; longer tenure lowers the risk.")
    rows = [["Model (test split)", "PR-AUC", "ROC-AUC", "Brier", "ECE"]]
    labels = [("dummy_raw", "Prior-only baseline"), (f"{m['selected']}_calibrated", "Logistic regression"),
              ("random_forest_calibrated", "Random forest"), ("xgboost_calibrated", "XGBoost")]
    for key, label in labels:
        r = test[key]
        rows.append([label, num(r["average_precision"], 3), num(r["roc_auc"], 3), num(r["brier"], 3), num(r["ece"], 3)])
    table(s, 0.5, 1.4, [3.3, 1.25, 1.25, 1.2, 1.2], rows, row_h=0.62, size=15, highlight=2)
    caption(s, 0.5, 4.6, 8.2, f"{num(f['data']['splits']['test'])} test customers held out from every choice; calibrated models; the baseline predicts the churn rate for everyone")
    for i, (value, label) in enumerate([(f"×{num(odds['Contract_Month-to-month'], 1)}", "odds: month-to-month vs two-year contract"),
                                        (f"×{num(odds['InternetService_Fiber optic'], 1)}", "odds: fibre vs no internet"),
                                        (f"×{num(odds['tenure'], 2)}", "odds per extra 12 months of tenure")]):
        stat(s, 0.5 + i * 2.8, 5.05, 2.6, 1.35, value, label, value_size=24, caption_size=12)
    stat(s, 9.1, 1.4, 3.75, 1.5, num(sel["roc_auc"], 2), f"ROC-AUC; PR-AUC {num(sel['average_precision'], 2)} vs {num(f['data']['churn_rate'], 2)} baseline")
    stat(s, 9.1, 3.1, 3.75, 1.5, num(sel["ece"], 3), "calibration error: probabilities fit to multiply by money")
    text(s, 9.1, 4.8, 3.75, 1.6, ["Logistic regression chosen by a rule fixed in advance: tree models did not rank better",
                                  "Drivers are associations, not causes"], size=14, bullets=True, space_after=4)

    # 8. Expected value and optimiser
    upgrade, bundle, standard = offers["contract_upgrade"], offers["service_bundle"], offers["standard"]
    value = f["cfg"]["value"]
    s = deck.slide("content", "From Probability to Offer: Expected Value", "Hieu", 55,
                   "Next we turn a probability into a decision. The expected value of an offer is the churn probability, times the save rate, "
                   "times the customer's value, minus the offer and contact cost. "
                   f"Customer value is the monthly bill times a {pct(value['margin'], 0)} margin over a discounted twelve-month horizon, longer for loyal customers and shorter for new ones. "
                   f"We defined a small catalogue of offers, for example a {usd(upgrade['offer_cost'])} credit for moving to a one-year contract, or six months of free tech support. "
                   "An integer program then chooses which customers get which offer to maximise total expected value within the budget, "
                   "with one offer per customer, eligibility rules and a contact cap. The save rates are assumptions, not measurements, and we treat them that way.")
    banner(s, 0.5, 1.4, 12.35, 1.0, ["EV  =  p(churn)  ×  save rate  ×  customer value  −  (offer cost + contact cost)"], size=24)
    head = {"size": 20, "bold": True, "color": NAVY}
    card(s, 0.5, 2.7, 3.95, 3.3, [[("Customer value", head)],
                                  f"Monthly bill × {pct(value['margin'], 0)} margin × discounted {value['base_horizon_months']} months",
                                  f"Horizon ×{value['tenure_factors']['new']} for new and ×{value['tenure_factors']['loyal']} for loyal customers",
                                  f"{pct(value['discount_rate_annual'], 0)} annual discount rate"], size=17, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space_after=10)
    card(s, 4.7, 2.7, 3.95, 3.3, [[("Offer catalogue (assumed)", head)],
                                  f"{upgrade['label']}: save rate {pct(upgrade['save_rate'], 0)}",
                                  f"{bundle['label']}: {pct(bundle['save_rate'], 0)}",
                                  f"{standard['label']}: {pct(standard['save_rate'], 0)}",
                                  f"Plus {usd(standard['contact_cost'])} contact cost per customer"], size=16, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space_after=8)
    card(s, 8.9, 2.7, 3.95, 3.3, [[("Optimiser", head)],
                                  "Integer program maximises total EV within the budget",
                                  "One offer per customer; eligibility rules",
                                  f"At most {pct(f['cfg']['campaign']['max_contact_share'], 0)} of customers contacted",
                                  "EV ≤ 0: never contacted"], size=17, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space_after=10)
    caption(s, 0.5, 6.15, 12.35, "Save rates and costs are visible assumptions, not measurements: the data contain no randomised campaign, so no causal uplift is claimed.")

    # 9. Results at the same budget
    unc = f["uncertainty"]
    s = deck.slide("content", f"Same Budget, {usd(hl['delta_vs_risk'])} More Expected Value", "Hieu", 50,
                   f"Here is the result on the {num(f['budget']['n_test'])} test customers, all with the same {usd(f['budget']['base'])} budget. "
                   "Calling random customers, or the most valuable ones, loses money, because most of them would have stayed anyway. "
                   f"Today's highest-risk rule earns about {usd(c['risk']['expected_net_value'])}. Expected-value targeting with segment offers earns {usd(chosen['expected_net_value'])}: "
                   f"that is {usd(hl['delta_vs_risk'])} more for the same spend, {usd(hl['targeting_effect'])} from choosing better customers and {usd(hl['offer_design_effect'])} from matching offers. "
                   "When we revalue the lists with customers' actual churn, the gain is still positive, "
                   f"with a 95% interval from {usd(hl['gain_label']['lower_95'])} to {usd(hl['gain_label']['upper_95'])}. Dung will now explain how much to spend and how sure we are.")
    bars = [("Random contact", c["random_mean"]["expected_net_value"]), ("Highest value", c["value"]["expected_net_value"]),
            ("Highest risk (today)", c["risk"]["expected_net_value"]), ("EV, one offer", c["ev"]["expected_net_value"]),
            ("EV + segment offers", chosen["expected_net_value"])]
    data = CategoryChartData()
    data.categories = [b[0] for b in bars]
    data.add_series("Expected net value", [round(b[1]) for b in bars])
    chart = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.5), Inches(1.35), Inches(7.6), Inches(4.75), data).chart
    chart.has_legend = False
    chart.has_title = True
    chart.chart_title.text_frame.text = f"Expected net value, same {usd(f['budget']['base'])} budget (test)"
    _font(chart.chart_title.text_frame.paragraphs[0].runs[0], 14, NAVY, True)
    plot = chart.plots[0]
    plot.gap_width = 55
    plot.has_data_labels = True
    labels_ = plot.data_labels
    labels_.number_format, labels_.number_format_is_linked = '$#,##0;-$#,##0', False
    labels_.position = XL_LABEL_POSITION.OUTSIDE_END
    labels_.font.size, labels_.font.bold, labels_.font.name = Pt(13), True, FONT
    labels_.font.color.rgb = INK
    series = plot.series[0]
    series.invert_if_negative = False  # otherwise PowerPoint draws negative bars white
    for i, (_, v) in enumerate(bars):
        point = series.points[i]
        point.format.fill.solid()
        point.format.fill.fore_color.rgb = RED if v < 0 else MUTED if i == 2 else NAVY
    for dPt in series._element.findall(qn("c:dPt")):
        # Each coloured point needs its own flag too; PowerPoint ignores the series-level one there.
        flag = dPt.makeelement(qn("c:invertIfNegative"), {"val": "0"})
        dPt.find(qn("c:idx")).addnext(flag)
    values = [v for _, v in bars]
    # Room on both sides so the outside labels stay inside the plot, clear of the category names.
    chart.value_axis.minimum_scale = min(0.0, min(values) * 1.9)
    chart.value_axis.maximum_scale = max(values) * 1.3
    chart.value_axis.visible = False
    chart.value_axis.has_major_gridlines = False
    axis = chart.category_axis
    axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    axis.tick_labels.font.size, axis.tick_labels.font.name = Pt(13), FONT
    axis.tick_labels.font.color.rgb = INK
    axis.format.line.color.rgb = MUTED
    caption(s, 0.5, 6.15, 7.6, "Expected values under the stated assumptions (USD, fictional sample company), not realised profit.")
    stat(s, 8.5, 1.4, 4.35, 1.4, f"+{usd(hl['delta_vs_risk'])}", f"vs the highest-risk rule; ROI {pct(chosen['roi'], 0)} vs {pct(c['risk']['roi'], 0)}")
    stat(s, 8.5, 3.0, 4.35, 1.4, f"+{usd(hl['gain_label']['label_based_simulated_net'])}",
         f"valued with actual churn; 95% interval {usd(hl['gain_label']['lower_95'])} to {usd(hl['gain_label']['upper_95'])}")
    profile = f["profile"]
    text(s, 8.5, 4.6, 4.35, 1.8, [f"Risk-only is more precise ({pct(c['risk']['precision_contacted'], 0)} vs {pct(chosen['precision_contacted'], 0)}) but spends on customers worth little",
                                  f"Targets: {pct(profile['by']['InternetService|Fiber optic']['list_share'], 0)} fibre, "
                                  f"{pct(profile['by']['Contract|Month-to-month']['list_share'], 0)} month-to-month, median bill {usd(profile['list']['median_monthly_charges'])}"],
         size=14, bullets=True, space_after=4)

    # 10. Outline: results and robustness
    deck.outline("Dung", 5, "Thanks, Hieu. Let me turn to the budget and to risk.", active=2)

    # 11. How much to spend
    rec = f["recommendation"]
    s = deck.slide("content", "How Much to Spend: The Profit Curve", "Dung", 40,
                   "How big should the budget be? Each extra step of budget returns less. Our rule, fixed on validation data, keeps adding budget only while "
                   f"the last step still returns at least {pct(rec['hurdle_marginal_roi'], 0)}. That gives enough budget for the standard offer to {pct(rec['recommended_share'], 0)} of customers. "
                   f"On the test customers this buys {num(reco['selected_count'])} offers for {usd(reco['cost'])}, with an expected net value of {usd(reco['expected_net_value'])}. "
                   f"For the whole base of {num(f['scaled']['base_population'])} customers that is about {usd(f['scaled']['recommended_budget'])}, "
                   f"or {usd(f['budget']['recommended_per_base_customer'], 2)} per customer. The rule stops a little before the point where no customer has a positive expected value, "
                   "because the last steps barely pay for themselves.")
    picture(s, FIGURES / "profit_curve.png", 0.5, 1.3, h=5.1)
    stat(s, 7.6, 1.4, 5.25, 1.45, pct(rec["recommended_share"], 0), f"budget share chosen on validation: the last step must still return ≥ {pct(rec['hurdle_marginal_roi'], 0)}")
    stat(s, 7.6, 3.05, 5.25, 1.45, usd(reco["expected_net_value"]), f"expected net value on test: {num(reco['selected_count'])} offers for {usd(reco['cost'])} (ROI {pct(reco['roi'], 0)})")
    text(s, 7.6, 4.7, 5.25, 1.7, [f"About {usd(f['scaled']['recommended_budget'])} for the {num(f['scaled']['base_population'])}-customer base "
                                  f"({usd(f['budget']['recommended_per_base_customer'], 2)} per customer)",
                                  f"Positive EV runs out at {num(f['curve']['positive_ev_customers'])} test customers ({usd(f['curve']['positive_ev_cost'])}); "
                                  "the hurdle stops earlier, where extra steps barely pay"], size=14, bullets=True, space_after=4)

    # 12. How sure are we
    s = deck.slide("content", "How Sure Are We?", "Dung", 55,
                   f"How sure are we? We varied every assumption one at a time, and then all of them at once in {num(f['monte_carlo']['draws'])} Monte Carlo draws. "
                   f"In {share(mc['base_vs_risk']['prob_gain'])} of the draws, expected-value targeting beats the risk-only rule: who to target is robust. "
                   f"How much an offer saves is not. The recommended list breaks even if the standard offer saves {pct(rbe_standard)} of would-be leavers, while we assumed {pct(q, 0)}, "
                   f"and it loses money in {pct(mc['recommended']['prob_loss'], 0)} of the draws. IBM's churn reasons add a warning: only {pct(f['reasons']['addressable_share'], 0)} "
                   "of churners left for price or a competitor, which an offer can address. Finally, our fairness audit shows that seniors are contacted "
                   f"{num(senior['ratio_to_reference'], 1)} times as often as other customers, partly because of higher risk and partly because of higher bills, "
                   "so every wave is reviewed.")
    for i, (v, label) in enumerate([(share(mc["base_vs_risk"]["prob_gain"]), f"of {num(f['monte_carlo']['draws'])} draws: EV targeting beats risk-only"),
                                    (pct(rbe_standard), f"break-even save rate of the standard offer for the recommended list (assumed {pct(q, 0)})"),
                                    (pct(mc["recommended"]["prob_loss"], 0), "chance the recommended list loses money"),
                                    (pct(f["reasons"]["addressable_share"], 0), "of churners left for price or a competitor: what an offer can address")]):
        stat(s, 0.5 + i * 3.15, 1.4, 2.9, 2.1, v, label, value_size=34, caption_size=14)
    banner(s, 0.5, 3.8, 12.35, 0.9, ["Who to target is robust; how much an offer saves is not, so measure it before scaling"], size=20)
    drivers = f["monte_carlo"]["drivers"]
    card(s, 0.5, 4.95, 6.05, 1.45, [[("Biggest sources of uncertainty", {"bold": True, "color": NAVY, "size": 16})],
                                    f"{drivers[0]['label']}, {drivers[1]['label'].lower()} and {drivers[2]['label'].lower()}"],
         size=14, align=PP_ALIGN.LEFT, space_after=3)
    card(s, 6.8, 4.95, 6.05, 1.45, [[("Fairness check", {"bold": True, "color": NAVY, "size": 16})],
                                    f"Seniors contacted {num(senior['ratio_to_reference'], 1)}× as often ({num(senior['risk_adjusted_ratio'], 1)}× after adjusting for risk); "
                                    "age and gender are not model inputs; every wave is reviewed"], size=14, align=PP_ALIGN.LEFT, space_after=3)

    # 13. Recommendation: pilot
    roadmap = f["cfg"]["project"]["roadmap"]
    s = deck.slide("content", "Recommendation: A Randomised Pilot Before Scaling", "Dung", 55,
                   "So our recommendation to the Head of Retention is to adopt expected-value targeting with segment offers, but to launch it as a randomised pilot. "
                   f"One wave of {num(pilot['list_size_per_wave'])} customers from the recommended list is split at random: half receive the offer, half are held out. "
                   f"After {pilot['outcome_window_days']} days we measure the real save rate of this offer mix, which we assume to be {pct(q_mix)}. "
                   f"The decision rule is agreed now: go if the effect is significant and at least the break-even of {pct(rbe)}; stop if it is clearly below. "
                   f"The offered half costs about {usd(pilot['wave_cost'])}, and if our assumption holds, "
                   f"the rule says go {pct(pilot['at_assumed']['p_go'], 0)} of the time. A decision is ready about {pilot['weeks_to_decision']} weeks after approval. "
                   f"{DEMO_SPEAKER} will now show our dashboard.")
    x = 0.5
    for i, phase in enumerate(roadmap):
        card(s, x, 1.45, 2.75, 1.55, [[(phase["phase"], {"size": 18, "bold": True, "color": NAVY})], [(f"{phase['weeks']} weeks", {"size": 16, "bold": True, "color": RED})],
                                      [(phase["owner"], {"size": 12})]], space_after=2)
        if i < len(roadmap) - 1:
            operator(s, x + 2.75, 1.45, 0.45, 1.55, "→")
        x += 3.2
    stat(s, 0.5, 3.3, 3.95, 1.6, num(pilot["list_size_per_wave"]), f"customers: {num(pilot['arm_size_per_wave'])} offered, {num(pilot['arm_size_per_wave'])} held out at random")
    stat(s, 4.7, 3.3, 3.95, 1.6, usd(pilot["wave_cost"]), f"cost of the offered half; expected net value {usd(pilot['wave_expected_net'])}")
    stat(s, 8.9, 3.3, 3.95, 1.6, f"≥ {pct(rbe)}", f"measured save rate of the offer mix needed to go (assumed {pct(q_mix)}); P(go) = {pct(pilot['at_assumed']['p_go'], 0)} if it holds")
    banner(s, 0.5, 5.2, 12.35, 1.15, ["Adopt expected-value targeting with segment offers, launched as a randomised pilot;",
                                      "scale only if the measured save rate clears break-even"], size=19)

    # 14. Live demo
    budget = f["budget"]["base"]
    s = deck.slide("content", "Live Demo: Retention Lab Dashboard", DEMO_SPEAKER, 0,
                   "Short version of reports/presentation/demo-script.md (Streamlit app, default assumptions, held-out test customers):\n"
                   f"1. Decision page: read the five KPIs (expected net value {usd(chosen['expected_net_value'])}, ROI {pct(chosen['roi'], 0)}, "
                   f"{num(chosen['selected_count'])} customers contacted, precision, high-risk/high-value coverage) and the green recommendation sentence.\n"
                   "2. Scroll to 'Who gets an offer': fibre, month-to-month and bill profile, then the risk x value scatter.\n"
                   f"3. Set the budget to 0 and back to {usd(budget)}: KPIs and the action list update at once.\n"
                   "4. Action list: search one customer ID, open 'Why is a customer in or out of the list?', then click 'Export campaign' (CSV + assumptions).\n"
                   "5. Assumptions & sensitivity: lower the save rate to 10%; the plan shrinks and the live chance of a loss rises; point at the break-even save rate.\n"
                   "6. Model evidence: performance table and points scorecard; upload data/examples/customers_unlabeled.csv to score new customers.\n"
                   "Hand back: 'That is the tool a retention manager would use for every wave. Back to Dung.'")
    if DASHBOARD.exists():
        picture(s, crop_dashboard(DASHBOARD), 0.5, 1.3, h=5.1)
    for i, (head_, body) in enumerate([("Decision", "Five KPIs, recommendation, who gets an offer, action list and export"),
                                        ("Assumptions & sensitivity", "Change save rates and costs; break-even, live chance of a loss, stress test"),
                                        ("Model evidence", "Performance, points scorecard, SHAP challenger, score an uploaded customer file")]):
        card(s, 8.2, 1.4 + i * 1.7, 4.65, 1.5, [[(head_, {"size": 18, "bold": True, "color": NAVY})], body], size=14, align=PP_ALIGN.LEFT, space_after=3)

    # 15. Thank you
    s = deck.slide("thanks", "Thank you", "Dung", 5, "Thank you for your attention. We are happy to take your questions.")
    text(s, 6.3, 3.0, 6.6, 1.1, ["THANK YOU !"], size=60, color=RGBColor.from_string("C00000"), bold=True)
    text(s, 6.35, 4.15, 6.5, 0.6, ["Questions and discussion"], size=24, color=NAVY)

    # 16. Closing cover
    deck.slide("closing", "Closing", None, 0, "Closing slide: leave it on screen during questions.")

    deck.save(out_path)
    if script_path is not None:
        write_script(deck.script, facts, script_path)
    return deck.script


def lint(script, min_rate=1.7, max_rate=2.6, balance_seconds=20):
    """Problems with the speaker script: missing notes, a talk that does not add up to its slot, an unbalanced split, or a pace too fast or slow to keep."""
    problems = [f"slide {e['slide']} has no speaker notes" for e in script if not e["text"].strip()]
    timed = [e for e in script if e["speaker"] in SPEAKERS]
    total = sum(e["seconds"] for e in timed)
    if total != TALK_SECONDS:
        problems.append(f"talk lasts {total} s; expected {TALK_SECONDS} s")
    for speaker in SPEAKERS:
        seconds = sum(e["seconds"] for e in timed if e["speaker"] == speaker)
        if abs(seconds - TALK_SECONDS / len(SPEAKERS)) > balance_seconds:
            problems.append(f"{speaker} speaks {seconds} s; expected about {TALK_SECONDS // len(SPEAKERS)} s")
    for e in timed:
        if e["seconds"] >= 30 and not min_rate <= len(e["text"].split()) / e["seconds"] <= max_rate:
            problems.append(f"slide {e['slide']}: {len(e['text'].split()) / e['seconds']:.1f} words per second")
    return problems


def _slide_ranges(numbers):
    """[10, 11, 12, 13, 15] -> '10-13, 15'."""
    groups, start = [], None
    for i, n in enumerate(numbers):
        start = n if start is None else start
        if i == len(numbers) - 1 or numbers[i + 1] != n + 1:
            groups.append(f"{start}-{n}" if n > start else str(n))
            start = None
    return ", ".join(groups)


def write_script(script, facts, path):
    """Rehearsal copy of the slide notes: timing per speaker, then the script slide by slide."""
    timed = [e for e in script if e["speaker"] in SPEAKERS]
    lines = [f"<!-- Generated by src/telco_retention/presentation.py for run {facts['run']['id']}. Edit the code, not this file. -->",
             "# Presentation script", "",
             f"{TALK_SECONDS // 60}-minute talk by {', '.join(SPEAKERS[:-1])} and {SPEAKERS[-1]}; {DEMO_SPEAKER} runs the live dashboard demo after the talk. "
             "The same text is in the speaker notes of `retention-presentation.pptx`.", "",
             "| Speaker | Slides | Time |", "|---|---|---:|"]
    for speaker in [*SPEAKERS, DEMO_SPEAKER]:
        part = [e for e in script if e["speaker"] == speaker]
        seconds = sum(e["seconds"] for e in part)
        time = f"{seconds // 60}:{seconds % 60:02d}" if speaker in SPEAKERS else "demo, not timed"
        lines.append(f"| {speaker} | {_slide_ranges([e['slide'] for e in part])} | {time} |")
    total = sum(e["seconds"] for e in timed)
    lines += [f"| **Total** | | **{total // 60}:{total % 60:02d}** |", ""]
    for e in script:
        lines += [f"## Slide {e['slide']}: {e['title']}", "", f"*{e['cue'].strip('[]')}*", "", e["text"], ""]
    Path(path).write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8", newline="\n")
    return Path(path)
