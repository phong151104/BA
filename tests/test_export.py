import copy
import pytest
from docx import Document
from pptx import Presentation
from telco_retention.common import ROOT, latest_run
from telco_retention.export import markdown_to_docx
from telco_retention.facts import build_facts
from telco_retention.presentation import DEMO_SPEAKER, SPEAKERS, TALK_SECONDS, _slide_ranges, build_presentation, lint
from telco_retention.templating import render


@pytest.fixture(scope="module")
def facts():
    return build_facts(latest_run())


@pytest.fixture(scope="module")
def deck(facts, tmp_path_factory):
    folder = tmp_path_factory.mktemp("presentation")
    script = build_presentation(facts, folder / "deck.pptx", folder / "script.md")
    return Presentation(folder / "deck.pptx"), script, (folder / "script.md").read_text(encoding="utf-8")


def test_brief_docx_has_structure(tmp_path):
    md = tmp_path / "brief.md"
    md.write_text(render("reports/executive-brief.md.j2", build_facts(latest_run())), encoding="utf-8")
    (tmp_path / "figures").mkdir()
    (tmp_path / "figures/profit_curve.png").write_bytes((ROOT / "reports/figures/profit_curve.png").read_bytes())
    doc = Document(markdown_to_docx(md, tmp_path / "brief.docx"))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert headings[0].startswith("Executive brief") and "Decisions requested" in headings
    assert len(doc.tables) == 1 and len(doc.inline_shapes) == 1
    assert all("**" not in p.text for p in doc.paragraphs)


def test_deck_has_a_speaker_script_on_every_slide(deck):
    prs, script, _ = deck
    assert lint(script) == []
    assert len(prs.slides) == len(script)
    for slide, entry in zip(prs.slides, script):
        notes = slide.notes_slide.notes_text_frame.text
        assert notes.startswith(entry["cue"]) and entry["text"] in notes


def test_talk_fits_the_slot_and_is_shared(deck):
    _, script, _ = deck
    timed = [e for e in script if e["speaker"] in SPEAKERS]
    assert sum(e["seconds"] for e in timed) == TALK_SECONDS
    assert {e["speaker"] for e in timed} == set(SPEAKERS)
    # Each speaker talks in one block, in order, and the demo follows the talk.
    order = [e["speaker"] for e in timed]
    assert order == sorted(order, key=SPEAKERS.index)
    demo = next(e for e in script if e["speaker"] == DEMO_SPEAKER)
    assert demo["slide"] > max(e["slide"] for e in timed if e["seconds"] >= 30)


def test_deck_uses_the_template_and_the_facts(deck, facts):
    prs, script, rehearsal = deck
    assert len(prs.slide_masters) == 5 and prs.slides[0].slide_layout.slide_master == prs.slide_masters[0]
    text = " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    gain = f"${facts['headline']['delta_vs_risk']:,.0f}"
    assert gain in text and gain in " ".join(e["text"] for e in script) and gain in rehearsal
    assert f"{facts['pilot']['list_size_per_wave']:,}" in text


def test_deck_follows_changed_facts(facts, tmp_path):
    changed = copy.deepcopy(facts)
    changed["headline"]["delta_vs_risk"] = 4321.0
    script = build_presentation(changed, tmp_path / "deck.pptx")
    assert any("$4,321" in e["text"] for e in script)


def test_lint_flags_timing_and_pace():
    script = [{"slide": 1, "speaker": "Kiet", "seconds": 300, "text": "word " * 900},
              {"slide": 2, "speaker": "Hieu", "seconds": 40, "text": ""},
              {"slide": 3, "speaker": "Dung", "seconds": 0, "text": "Thanks."}]
    problems = lint(script)
    assert any("talk lasts" in p for p in problems) and any("slide 2 has no speaker notes" in p for p in problems)
    assert "slide 1: 3.0 words per second" in problems  # too fast
    assert "slide 2: 0.0 words per second" in problems  # silent
    assert any(p.startswith("Dung speaks") for p in problems)


def test_deck_avoids_overclaims(deck):
    _, script, rehearsal = deck
    banned = ["the method works", "never loses money", "causal uplift of", "guaranteed", "proves"]
    lowered = rehearsal.lower()
    assert not [b for b in banned if b in lowered]


def test_slide_ranges():
    assert _slide_ranges([10, 11, 12, 13, 15]) == "10-13, 15"
    assert _slide_ranges([14]) == "14"
